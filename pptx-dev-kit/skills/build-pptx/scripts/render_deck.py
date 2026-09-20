#!/usr/bin/env python3
"""Render deck.json to a .pptx using the kit's layout catalog. Do not replace this with an ad-hoc script."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from engine.chrome import add_chrome  # noqa: E402
from engine.grid import SW, SH, emu_length  # noqa: E402
from engine.shapes import add_notes, drop_empty_placeholders  # noqa: E402
from engine.theme import Theme  # noqa: E402
from layouts import LAYOUTS  # noqa: E402
from validate_deck_json import validate_deck  # noqa: E402


def blank_layout(prs):
    for layout in prs.slide_layouts:
        if (layout.name or "").strip().lower() == "blank":
            return layout
    return prs.slide_layouts[6]


def render(doc: dict, out_path: Path) -> None:
    errors = validate_deck(doc)
    if errors:
        raise SystemExit("deck.json invalid:\n" + "\n".join(f"  - {e}" for e in errors))
    theme = Theme(doc["theme"])
    prs = Presentation()
    prs.slide_width = Emu(emu_length(SW))
    prs.slide_height = Emu(emu_length(SH))
    layout = blank_layout(prs)
    slides = doc["slides"]
    for i, spec in enumerate(slides, 1):
        name = spec["layout"]
        painter = LAYOUTS.get(name)
        if painter is None:
            raise SystemExit(f"unknown layout {name!r}")
        slide = prs.slides.add_slide(layout)
        painter(slide, spec, theme, i)
        add_chrome(slide, theme, i, name)
        add_notes(slide, spec.get("notes"))
        drop_empty_placeholders(slide)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    check = Presentation(str(out_path))
    n = len(check.slides._sldIdLst)
    if n != len(slides):
        raise SystemExit(f"re-open slide count {n} != {len(slides)}")
    print(f"OK: {out_path} ({n} slides)")


def main() -> int:
    parser = argparse.ArgumentParser(description="Render pptx-dev-kit deck.json")
    parser.add_argument("deck", help="Path to deck.json")
    parser.add_argument("-o", "--output", required=True, help="Output .pptx path")
    args = parser.parse_args()
    path = Path(args.deck)
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        print(f"ERROR: file not found: {path}", file=sys.stderr)
        return 2
    except json.JSONDecodeError as exc:
        print(f"ERROR: invalid JSON: {exc}", file=sys.stderr)
        return 1
    try:
        render(doc, Path(args.output))
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
