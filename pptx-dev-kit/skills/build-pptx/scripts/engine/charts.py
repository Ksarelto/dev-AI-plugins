"""Native 2-D charts with palette series colors and required alt text."""

from __future__ import annotations

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu

from .colors import rgb
from .grid import emu_length
from .shapes import set_alt_text
from .theme import Theme


def _in(inches: float):
    return Emu(emu_length(inches))

_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def add_chart(slide, spec: dict, left, top, width, height, *, theme: Theme):
    chart_type = (spec.get("type") or "column").lower()
    if chart_type not in _TYPES:
        raise ValueError(f"unsupported chart type {chart_type!r}; use {sorted(_TYPES)}")
    if chart_type in {"pie", "doughnut"} and len(spec.get("series") or []) != 1:
        raise ValueError(f"{chart_type} charts support exactly one series")

    data = CategoryChartData()
    data.categories = list(spec.get("categories") or [])
    for series in spec.get("series") or []:
        data.add_series(series["name"], tuple(series["values"]))

    gf = slide.shapes.add_chart(
        _TYPES[chart_type],
        _in(left),
        _in(top),
        _in(width),
        _in(height),
        data,
    )
    chart = gf.chart
    series_list = list(spec.get("series") or [])
    multi = len(series_list) > 1 and chart_type not in {"pie", "doughnut"}
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    plot = chart.plots[0]
    plot.has_data_labels = True
    fmt = spec.get("number_format") or "0"
    plot.data_labels.number_format = fmt
    try:
        plot.data_labels.font.size = None
    except Exception:
        pass

    palette = [
        theme.c("accent"),
        theme.colors.get("text", theme.c("text")),
        theme.colors.get("muted", theme.c("muted")),
        theme.c("cover_bg"),
    ]
    if chart_type in {"pie", "doughnut"}:
        series = chart.series[0]
        for i, _cat in enumerate(spec.get("categories") or []):
            try:
                point = series.points[i]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
            except Exception:
                break
    else:
        for i, series in enumerate(chart.series):
            color = rgb(palette[i % len(palette)])
            try:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = color
            except Exception:
                pass
            try:
                series.format.line.color.rgb = color
            except Exception:
                pass

    try:
        plot.has_data_labels = True
        if hasattr(chart, "value_axis") and chart.value_axis is not None:
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = rgb(theme.c("muted"))
    except Exception:
        pass

    alt = spec.get("alt") or "Chart"
    set_alt_text(gf, alt)
    return gf
