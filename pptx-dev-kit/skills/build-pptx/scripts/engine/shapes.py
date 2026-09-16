"""Shape and text primitives used by every layout recipe."""

from __future__ import annotations

import uuid

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .colors import rgb
from .grid import emu_length
from .theme import Theme


def _in(inches: float):
    """Inches rounded to a whole EMU — Inches() truncates and lets gutters drift."""
    return Emu(emu_length(inches))


def add_soft_shadow(shape) -> None:
    """Subtle depth under cards (2026 soft-curve / layered-module look). Skip on tiny bars."""
    sp_pr = shape._element.spPr
    for child in list(sp_pr):
        if child.tag == qn("a:effectLst"):
            sp_pr.remove(child)
    effect = etree.SubElement(sp_pr, qn("a:effectLst"))
    shadow = etree.SubElement(
        effect,
        qn("a:outerShdw"),
        blurRad="50800",
        dist="19050",
        dir="2700000",
        algn="tl",
        rotWithShape="0",
    )
    color = etree.SubElement(shadow, qn("a:srgbClr"), val="1B1B1B")
    etree.SubElement(color, qn("a:alpha"), val="16000")


_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}

_ANCHOR = {
    "top": "t",
    "middle": "ctr",
    "bottom": "b",
}


def fill_slide(slide, hex_color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)


def add_rect(slide, left, top, width, height, fill_hex, *, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, _in(left), _in(top), _in(width), _in(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_hex)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
    return shape


def add_round_rect(
    slide, left, top, width, height, fill_hex, *, radius=0.12, line=None, shadow=False
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        _in(left),
        _in(top),
        _in(width),
        _in(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_hex)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
    try:
        shape.adjustments[0] = radius
    except (AttributeError, IndexError, ValueError):
        pass
    if shadow:
        add_soft_shadow(shape)
    return shape


def add_oval(slide, left, top, width, height, fill_hex):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, _in(left), _in(top), _in(width), _in(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_hex)
    shape.line.fill.background()
    return shape


def _apply_run_style(run, *, font, size_pt, color, bold=False, italic=False):
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold
    run.font.italic = italic


def add_text(
    slide,
    text,
    left,
    top,
    width,
    height,
    *,
    theme: Theme,
    size_pt,
    color,
    font=None,
    bold=False,
    italic=False,
    align="left",
    anchor="top",
    wrap=True,
):
    tb = slide.shapes.add_textbox(_in(left), _in(top), _in(width), _in(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.clear()
    body_pr = tf._txBody.find(qn("a:bodyPr"))
    if body_pr is not None:
        body_pr.set("anchor", _ANCHOR.get(anchor, "t"))
    p = tf.paragraphs[0]
    p.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = text or ""
    _apply_run_style(
        run,
        font=font or theme.body_font,
        size_pt=size_pt,
        color=color,
        bold=bold,
        italic=italic,
    )
    return tb


def add_lines(
    slide,
    lines,
    left,
    top,
    width,
    height,
    *,
    theme: Theme,
    size_pt,
    color,
    font=None,
    bold=False,
    align="left",
    space_after_pt=8,
):
    tb = slide.shapes.add_textbox(_in(left), _in(top), _in(width), _in(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    tf.clear()
    if not lines:
        lines = [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
        p.space_after = Pt(space_after_pt)
        run = p.add_run()
        run.text = line
        _apply_run_style(
            run,
            font=font or theme.body_font,
            size_pt=size_pt,
            color=color,
            bold=bold,
        )
    return tb


def add_bullet_rows(slide, items, left, top, width, *, theme: Theme, size_pt=None, row_h=0.42):
    """Accent disc + text, vertically centered on each row so marks don't drift."""
    size_pt = size_pt or theme.pt_body
    mark = 0.14
    y = top
    for item in items:
        add_oval(slide, left, y + (row_h - mark) / 2, mark, mark, theme.c("accent"))
        add_text(
            slide,
            item,
            left + mark + 0.14,
            y,
            width - mark - 0.14,
            row_h,
            theme=theme,
            size_pt=size_pt,
            color=theme.c("text"),
            font=theme.body_font,
            anchor="middle",
        )
        y += row_h
    return y


def add_spark_bars(
    slide,
    values,
    left,
    top,
    width,
    height,
    *,
    theme: Theme,
    accent_last=True,
    color=None,
    rest_color=None,
):
    """Tiny column bars for a KPI/chart card. Values are relative magnitudes."""
    if not values:
        return
    cleaned = [max(float(v), 0.0) for v in values]
    peak = max(cleaned) or 1.0
    n = len(cleaned)
    gap = min(0.06, width / max(n * 4, 1))
    bar_w = (width - gap * (n - 1)) / n
    color = color or theme.c("accent")
    rest_color = rest_color or theme.c("muted")
    for i, value in enumerate(cleaned):
        h = max(height * (value / peak), 0.08)
        x = left + i * (bar_w + gap)
        y = top + height - h
        fill = color if (accent_last and i == n - 1) else rest_color
        add_round_rect(slide, x, y, bar_w, h, fill, radius=0.08)


def add_disc_label(slide, left, top, size, text, *, theme: Theme, fill=None):
    fill = fill or theme.c("accent")
    add_oval(slide, left, top, size, size, fill)
    add_text(
        slide,
        text,
        left,
        top,
        size,
        size,
        theme=theme,
        size_pt=11 if size < 0.42 else 13,
        color=theme.c("accent_text"),
        font=theme.body_font,
        bold=True,
        align="center",
        anchor="middle",
    )


def set_alt_text(shape, text: str) -> None:
    """Write descr on the nvPr element. `shape.alt_text = ...` is a no-op."""
    if not text:
        return
    shape._element._nvXxPr.cNvPr.set("descr", text)


def drop_empty_placeholders(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def add_notes(slide, notes) -> None:
    if not notes:
        return
    if isinstance(notes, str):
        text = notes
    else:
        text = "\n".join(n for n in notes if n)
    if text.strip():
        slide.notes_slide.notes_text_frame.text = text


_FLD = (
    '<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'id="{{{id}}}" type="slidenum">'
    '<a:rPr lang="en-US" sz="1100"/>'
    "<a:t>1</a:t></a:fld>"
)


def add_slide_number(slide, left, top, width, height, *, theme: Theme, color: str):
    guid = str(uuid.uuid4()).upper()
    tb = slide.shapes.add_textbox(_in(left), _in(top), _in(width), _in(height))
    tf = tb.text_frame
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p._p.append(etree.fromstring(_FLD.format(id=guid)))
    for el in p._p.iter():
        if el.tag == qn("a:rPr"):
            el.set("sz", "1100")
    return tb
