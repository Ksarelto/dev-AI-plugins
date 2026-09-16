#!/usr/bin/env python3
"""Plain-text extract of a .pptx when markitdown is not installed."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def extract(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks = []
    for i, slide in enumerate(prs.slides, 1):
        lines = [f"## Slide {i}"]
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        notes = ""
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            lines.append("Notes: " + notes.replace("\n", " / "))
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx")
    parser.add_argument("-o", "--output", default=None)
    args = parser.parse_args()
    path = Path(args.pptx)
    if not path.is_file():
        print(f"ERROR: not found: {path}", file=sys.stderr)
        return 2
    try:
        text = extract(path)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"OK: wrote {args.output}")
    else:
        sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
