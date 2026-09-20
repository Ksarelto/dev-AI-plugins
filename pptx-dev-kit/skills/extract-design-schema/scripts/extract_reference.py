#!/usr/bin/env python3
"""Extract the design system from a reference .pptx into JSON.

Deterministic extraction only — no interpretation. The `extract-design-schema` skill turns this
JSON into a design-schema.md; this script's only job is to report what is *actually in the file*,
including what it could not find.

Usage:
    python3 extract_reference.py REFERENCE.pptx [-o report.json] [--full]

Compact by default: layouts are summarized to their content-bearing placeholders, because a full
placeholder dump of 11 layouts is thousands of tokens of mostly-identical footer geometry. Pass
--full only when you need to reproduce an exact layout box.

Requires: python-pptx (brings lxml and Pillow).
"""
from __future__ import annotations

import argparse
import collections
import json
import os
import sys
import zipfile

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
EMU_PER_INCH = 914400

# dk1/lt1 are usually system colors; these are their standard resolutions.
SYS_COLOR = {"windowText": "000000", "window": "FFFFFF"}


def _in(emu):
    return round(emu / EMU_PER_INCH, 3) if emu is not None else None


def contrast_ratio(h1: str, h2: str) -> float:
    def lum(h):
        h = h.lstrip("#")
        c = [int(h[i : i + 2], 16) / 255 for i in (0, 2, 4)]
        c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
        return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]

    a, b = sorted((lum(h1), lum(h2)), reverse=True)
    return round((a + 0.05) / (b + 0.05), 2)


def extract_theme(zf: zipfile.ZipFile) -> dict:
    """Theme colors and fonts, from the first theme part."""
    from lxml import etree

    names = [n for n in zf.namelist() if n.startswith("ppt/theme/theme")]
    if not names:
        return {"error": "no theme part found"}
    root = etree.fromstring(zf.read(sorted(names)[0]))

    colors = {}
    scheme = root.find(".//a:clrScheme", NS)
    if scheme is not None:
        for node in scheme:
            tag = etree.QName(node).localname
            child = node[0] if len(node) else None
            if child is None:
                continue
            val = child.get("val")
            if val in SYS_COLOR:  # <a:sysClr val="windowText" lastClr="000000"/>
                val = child.get("lastClr") or SYS_COLOR[val]
            colors[tag] = (val or "").upper() or None

    fonts = {}
    fscheme = root.find(".//a:fontScheme", NS)
    if fscheme is not None:
        for key, which in (("heading", "a:majorFont"), ("body", "a:minorFont")):
            node = fscheme.find(f"{which}/a:latin", NS)
            if node is not None:
                fonts[key] = node.get("typeface") or None

    return {"colors": colors, "fonts": fonts}


# Placeholder types that carry content — the rest is chrome repeated on every layout.
CONTENT_PH = {"TITLE", "CENTER_TITLE", "SUBTITLE", "BODY", "OBJECT", "PICTURE", "CHART", "TABLE"}


def extract_geometry(prs, full: bool = False) -> dict:
    """Canvas size and the placeholder grid implied by the layouts."""
    w, h = prs.slide_width, prs.slide_height
    ratio = round(w / h, 3) if h else None
    aspect = {1.778: "16:9", 1.333: "4:3", 1.6: "16:10"}.get(round(ratio, 3), f"{ratio}:1")

    lefts, rights, tops = [], [], []
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        phs = []
        for ph in layout.placeholders:
            fmt = ph.placeholder_format
            kind = str(fmt.type).split(" ")[0]
            box = {
                "idx": fmt.idx,
                "type": kind,
                "left_in": _in(ph.left),
                "top_in": _in(ph.top),
                "width_in": _in(ph.width),
                "height_in": _in(ph.height),
            }
            if full or kind in CONTENT_PH:
                phs.append(box)
            if kind in CONTENT_PH and ph.left is not None and ph.width is not None:
                lefts.append(ph.left)
                rights.append(w - (ph.left + ph.width))
                if ph.top is not None:
                    tops.append(ph.top)
        layouts.append({"index": i, "name": layout.name, "placeholders": phs})

    margins = None
    if lefts:
        left_in, right_in = _in(min(lefts)), _in(min(rights))
        margins = {
            "left_in": left_in,
            "right_in": right_in,
            "top_in": _in(min(tops)) if tops else None,
            "symmetric": abs(left_in - right_in) < 0.15,
        }

    return {
        "slide_width_in": _in(w),
        "slide_height_in": _in(h),
        "aspect_ratio": aspect,
        "implied_margins": margins,
        "layouts": layouts,
    }


def extract_usage(prs) -> dict:
    """What the deck ACTUALLY uses on its slides — the theme only says what was declared."""
    sizes = collections.Counter()
    fonts = collections.Counter()
    colors = collections.Counter()
    title_sizes, body_sizes = [], []
    charts, pictures, tables = 0, 0, 0
    notes_slides = 0

    for slide in prs.slides:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes_slides += 1
        for shape in slide.shapes:
            if shape.has_chart:
                charts += 1
            if shape.shape_type == 13:
                pictures += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if not shape.has_text_frame:
                continue
            is_title = shape == slide.shapes.title
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    f = run.font
                    if f.size is not None:
                        pt = f.size.pt
                        sizes[pt] += 1
                        (title_sizes if is_title else body_sizes).append(pt)
                    if f.name:
                        fonts[f.name] += 1
                    try:
                        if f.color is not None and f.color.type is not None and f.color.rgb is not None:
                            colors[str(f.color.rgb).upper()] += 1
                    except (AttributeError, TypeError, ValueError):
                        pass  # theme-colored runs have no direct rgb

    return {
        "slide_count": len(prs.slides._sldIdLst),
        "explicit_font_sizes_pt": dict(sizes.most_common()),
        "explicit_fonts": dict(fonts.most_common()),
        "explicit_run_colors": dict(colors.most_common()),
        "title_sizes_pt": sorted(set(title_sizes)),
        "body_sizes_pt": sorted(set(body_sizes)),
        "charts": charts,
        "pictures": pictures,
        "tables": tables,
        "slides_with_speaker_notes": notes_slides,
        "note": (
            "Empty size/font/color maps mean the deck inherits everything from its master — that is "
            "a well-built deck, not a failed extraction. Read the theme block for the real values."
        ),
    }


def derive_pairs(theme: dict) -> list:
    """Contrast ratios for the pairings a deck actually renders."""
    colors = theme.get("colors") or {}
    out = []
    for fg, bg, role in (
        ("dk1", "lt1", "body text on light background"),
        ("lt1", "dk1", "light text on dark background"),
        ("dk1", "accent1", "text on accent block"),
        ("lt1", "accent1", "light text on accent block"),
        ("dk2", "lt2", "secondary text on secondary background"),
    ):
        a, b = colors.get(fg), colors.get(bg)
        if not (a and b and len(a) == 6 and len(b) == 6):
            continue
        r = contrast_ratio(a, b)
        out.append(
            {
                "role": role,
                "foreground": f"#{a}",
                "background": f"#{b}",
                "ratio": r,
                "passes_AA_body": r >= 4.5,
                "passes_AA_large": r >= 3.0,
            }
        )
    return out


def signals(report: dict) -> list:
    """Things the skill must not silently inherit from the reference."""
    out = []
    theme, geo, usage = report["theme"], report["geometry"], report["usage"]

    if theme.get("error"):
        out.append("No theme part — palette and fonts could not be read; ask the user for brand values.")

    for pair in report["contrast_pairs"]:
        if not pair["passes_AA_large"]:
            out.append(
                f"Reference FAILS WCAG AA: {pair['role']} ({pair['foreground']} on "
                f"{pair['background']}) = {pair['ratio']}:1. Do not copy this pairing — "
                "substitute a compliant foreground and say so."
            )

    m = geo.get("implied_margins") or {}
    if m and m.get("symmetric") is False:
        out.append(
            f"Asymmetric margins (left {m['left_in']}in vs right {m['right_in']}in) — the layouts "
            f"were probably built for a different canvas than this {geo['aspect_ratio']} one. "
            "Derive the grid from the smaller margin rather than copying placeholder boxes."
        )

    body = usage.get("body_sizes_pt") or []
    if body and min(body) < 24:
        out.append(
            f"Reference uses body text at {min(body)}pt, below this kit's 24pt minimum. "
            "Keep the reference's typeface, not its size."
        )
    titles = usage.get("title_sizes_pt") or []
    if titles and min(titles) < 36:
        out.append(f"Reference uses titles at {min(titles)}pt, below this kit's 36pt minimum.")

    if not usage.get("explicit_fonts") and not (theme.get("fonts") or {}).get("body"):
        out.append("No font information found at theme or slide level — ask the user.")

    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--full", action="store_true", help="include every placeholder, not just content ones")
    args = ap.parse_args()

    try:
        from pptx import Presentation
    except ImportError:
        print("python-pptx is not installed: python3 -m pip install python-pptx", file=sys.stderr)
        return 2

    if not os.path.isfile(args.pptx):
        print(f"no such file: {args.pptx}", file=sys.stderr)
        return 2

    with open(args.pptx, "rb") as fh:
        magic = fh.read(8)
    if magic.startswith(b"\xd0\xcf\x11\xe0"):  # OLE2 compound file
        print(
            f"{args.pptx} is a legacy .ppt (or an encrypted Office file), not a .pptx. "
            "Open it in PowerPoint and Save As .pptx, then re-run.",
            file=sys.stderr,
        )
        return 2
    if not magic.startswith(b"PK"):
        print(f"{args.pptx} is not a .pptx — expected a zip container.", file=sys.stderr)
        return 2

    try:
        prs = Presentation(args.pptx)
        zf = zipfile.ZipFile(args.pptx)
    except Exception as exc:  # noqa: BLE001 - report verbatim; the caller needs the real reason
        print(f"could not open {args.pptx}: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    theme = extract_theme(zf)
    report = {
        "source": args.pptx,
        "theme": theme,
        "geometry": extract_geometry(prs, full=args.full),
        "usage": extract_usage(prs),
        "contrast_pairs": derive_pairs(theme),
    }
    report["signals"] = signals(report)

    text = json.dumps(report, indent=2)
    if args.out:
        with open(args.out, "w") as fh:
            fh.write(text)
        print(f"OK: wrote {args.out}")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
