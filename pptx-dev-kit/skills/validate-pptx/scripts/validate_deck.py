#!/usr/bin/env python3
"""
validate_deck.py — pptx-dev-kit post-build validator

Usage:
    python3 validate_deck.py <deck.pptx> [--expected-slides N] [--allow-ratio]

Exit codes:
    0  all checks pass (warnings allowed)
    1  one or more ERROR-level failures
    2  file not found or not a valid .pptx
"""
from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

PLACEHOLDER_OK = "[PLACEHOLDER — replace with actual data]"
GHOST_RE = re.compile(r"click to add", re.I)
LOREM_RE = re.compile(r"\blorem\b|\bipsum\b", re.I)
PLACEHOLDER_RE = re.compile(r"\bPLACEHOLDER\b")


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text or ""


def _leftover_placeholder(text: str) -> bool:
    cleaned = text.replace(PLACEHOLDER_OK, "")
    if GHOST_RE.search(cleaned) or LOREM_RE.search(cleaned):
        return True
    return bool(PLACEHOLDER_RE.search(cleaned))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--expected-slides", type=int, default=None)
    parser.add_argument(
        "--allow-ratio",
        action="store_true",
        help="Do not error when the canvas is not 16:9 (edit mode keeping source geometry)",
    )
    args = parser.parse_args()
    path = args.path
    errors: list[str] = []
    warnings: list[str] = []

    try:
        with open(path, "rb") as f:
            magic = f.read(4)
        if magic != b"PK\x03\x04":
            print(f"ERROR: {path} is not a valid .pptx (not a zip/PK file)", file=sys.stderr)
            sys.exit(2)
    except FileNotFoundError:
        print(f"ERROR: file not found: {path}", file=sys.stderr)
        sys.exit(2)

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches
    except ImportError:
        print("ERROR: python-pptx not installed", file=sys.stderr)
        sys.exit(2)

    try:
        prs = Presentation(path)
    except Exception as exc:
        print(f"ERROR: python-pptx could not open the file: {exc}", file=sys.stderr)
        sys.exit(2)

    SW, SH = prs.slide_width, prs.slide_height
    ratio = SW / SH if SH else 0
    is_169 = abs(ratio - 16 / 9) < 0.01
    is_43 = abs(ratio - 4 / 3) < 0.01
    if not is_169:
        msg = f"Canvas ratio {ratio:.3f} is not 16:9 ({SW/914400:.3f}in × {SH/914400:.3f}in)"
        if args.allow_ratio:
            warnings.append(msg + " — allowed by --allow-ratio")
        elif is_43:
            errors.append(msg + "; default template 4:3 leaked through (pass --allow-ratio to keep)")
        else:
            errors.append(msg + " (pass --allow-ratio to keep)")

    slide_count = len(prs.slides._sldIdLst)
    if args.expected_slides is not None and slide_count != args.expected_slides:
        errors.append(
            f"Slide count mismatch: got {slide_count}, expected {args.expected_slides}"
        )
    print(f"Slides: {slide_count}")

    inch = 914400
    title_band = int(Inches(1.30))
    engine_root = Path(__file__).resolve().parents[2] / "build-pptx" / "scripts"
    if str(engine_root) not in sys.path:
        sys.path.insert(0, str(engine_root))
    try:
        from engine.grid import ALIGN_TOL_IN, column_lefts, emu_length
    except ImportError:
        column_lefts = None
        ALIGN_TOL_IN = 0.02
        emu_length = lambda v: int(round(v * inch))  # noqa: E731
    rail_emus = None
    if column_lefts is not None:
        rail_emus = {0, *[emu_length(v) for v in column_lefts()]}

    for i, slide in enumerate(prs.slides, 1):
        for ph in slide.placeholders:
            if ph.has_text_frame and not ph.text_frame.text.strip():
                errors.append(
                    f"Slide {i}: unfilled placeholder idx={ph.placeholder_format.idx} "
                    f"('{ph.name}') — will render as ghost in PowerPoint"
                )

        for sh in slide.shapes:
            text = _shape_text(sh)
            if text and _leftover_placeholder(text):
                errors.append(
                    f"Slide {i}: leftover placeholder copy in '{sh.name}': {text[:80]!r}"
                )

            if getattr(sh, "has_text_frame", False):
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        size = run.font.size
                        if size is None:
                            continue
                        pt = size.pt
                        if pt < 11:
                            errors.append(
                                f"Slide {i}: run {pt:.1f}pt is below the 11pt caption floor "
                                f"in '{sh.name}'"
                            )
                        top = sh.top or 0
                        if top < title_band and 24 <= pt < 36:
                            errors.append(
                                f"Slide {i}: title-band run {pt:.1f}pt < 36pt in '{sh.name}'"
                            )

            if sh.shape_type == 13:  # PICTURE
                descr = sh._element._nvXxPr.cNvPr.get("descr", "")
                if not descr:
                    errors.append(
                        f"Slide {i}: picture '{sh.name}' has no alt text "
                        f"(descr attribute empty or missing)"
                    )
            if getattr(sh, "has_chart", False):
                descr = sh._element._nvXxPr.cNvPr.get("descr", "")
                if not descr:
                    errors.append(
                        f"Slide {i}: chart '{sh.name}' has no alt text "
                        f"(descr attribute empty or missing)"
                    )

            if sh.left is None or sh.top is None:
                continue
            if sh.left < -inch or sh.top < -inch:
                errors.append(
                    f"Slide {i}: shape '{sh.name}' positioned off canvas "
                    f"(left={sh.left/inch:.2f}in top={sh.top/inch:.2f}in)"
                )
            right_edge = sh.left + (sh.width or 0)
            bottom_edge = sh.top + (sh.height or 0)
            if right_edge > SW + inch:
                errors.append(
                    f"Slide {i}: shape '{sh.name}' overflows right edge "
                    f"(right={right_edge/inch:.2f}in, canvas={SW/inch:.2f}in)"
                )
            if bottom_edge > SH + inch:
                errors.append(
                    f"Slide {i}: shape '{sh.name}' overflows bottom edge "
                    f"(bottom={bottom_edge/inch:.2f}in, canvas={SH/inch:.2f}in)"
                )

            try:
                is_round_card = (
                    sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
                    and (sh.width or 0) >= 2.8 * inch
                    and (sh.height or 0) >= 0.9 * inch
                )
            except (ValueError, AttributeError):
                is_round_card = False
            if rail_emus and is_round_card:
                nearest = min(rail_emus, key=lambda r: abs(r - sh.left))
                if abs(sh.left - nearest) > ALIGN_TOL_IN * inch:
                    errors.append(
                        f"Slide {i}: card '{sh.name}' left={sh.left/inch:.3f}in "
                        f"is off the 12-column rail (nearest {nearest/inch:.3f}in)"
                    )

        notes_text = ""
        try:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes_text = ""
        if i > 1 and not notes_text:
            warnings.append(f"Slide {i}: no speaker notes")

    with zipfile.ZipFile(path) as zf:
        all_guids = []
        for name in sorted(zf.namelist()):
            if re.match(r"ppt/slides/slide\d+\.xml$", name):
                xml = zf.read(name).decode("utf-8", errors="replace")
                found = re.findall(
                    r'type="slidenum"[^>]*id="\{([^}]+)\}"', xml
                ) + re.findall(
                    r'id="\{([^}]+)\}"[^>]*type="slidenum"', xml
                )
                all_guids.extend(found)
        if len(all_guids) != len(set(all_guids)):
            from collections import Counter

            dupes = [g for g, c in Counter(all_guids).items() if c > 1]
            errors.append(
                f"Duplicate slide-number field GUIDs {dupes} — "
                f"all slides will show the same number in PowerPoint"
            )
        elif all_guids:
            print(f"Slide-number GUIDs: {len(all_guids)} unique")

    try:
        from lxml import etree as ET

        with zipfile.ZipFile(path) as zf:
            for name in zf.namelist():
                if name.startswith("ppt/") and name.endswith(".xml"):
                    try:
                        ET.fromstring(zf.read(name))
                    except ET.XMLSyntaxError as exc:
                        errors.append(f"Malformed XML in {name}: {exc}")
    except ImportError:
        warnings.append("lxml not available — XML well-formedness check skipped")

    print()
    if errors:
        for e in errors:
            print(f"ERROR: {e}")
    if warnings:
        for w in warnings:
            print(f"WARN:  {w}")

    if not errors and not warnings:
        print("All checks passed.")
    elif not errors:
        print(f"\nPASS with {len(warnings)} warning(s).")
    else:
        print(f"\nFAIL: {len(errors)} error(s), {len(warnings)} warning(s).")
        sys.exit(1)


if __name__ == "__main__":
    main()
